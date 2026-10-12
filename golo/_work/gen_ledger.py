# -*- coding: utf-8 -*-
"""生成高乐捷内容复盘两张台账 Excel 工作簿"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
import os

# 台账属于运营件，输出到 04-运营/（脚本迁移自旧目录 D:\Projects\golovejoy\复盘\）
OUT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '04-运营', '高乐捷复盘台账.xlsx'))

FONT_NAME = "微软雅黑"
HDR_FILL = PatternFill("solid", fgColor="1F3864")   # 深蓝
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")   # 浅蓝
HDR_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)
EX_FONT = Font(name=FONT_NAME, size=10, color="808080", italic=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = openpyxl.Workbook()

# ============ 表A：测试标记台账 ============
wsA = wb.active
wsA.title = "台账A-测试标记"

colsA = [
    ("素材ID", 12), ("素材名称/链接", 22), ("SPU", 10), ("素材类型", 12),
    ("测试假设(一句话,必填)", 30), ("主卖点标签", 14), ("目标人群标签", 14),
    ("创作责任人", 10), ("发布账号", 10), ("投放计划ID", 12),
    ("上线日期", 11), ("状态", 10), ("消耗(元)", 10), ("曝光", 10),
    ("点击率%", 9), ("转化率%", 9), ("ROI", 8),
    ("高光帧位置/内容", 18), ("结论", 12), ("评论区反馈摘要", 20), ("备注", 12),
]
wsA.append([c[0] for c in colsA])
for i, (_, w) in enumerate(colsA, 1):
    wsA.column_dimensions[get_column_letter(i)].width = w
    cell = wsA.cell(row=1, column=i)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER

exampleA = [
    "S-2026-0131", "https://v.douyin.com/xxxx", "保暖背心-旗舰", "翻拍",
    "验证:收腹卖点放开头前3秒能否提高点击率", "收腹显瘦", "微胖大胸",
    "佳敏", "主号", "QJ1024-07", "2026-10-07", "已出数据",
    12500, 93000, 3.2, 2.1, 3.5,
    "第4秒·收腹对比画面", "验证成功", "会卷边(2人3条评论)", "市场验证过的爆款脚本",
]
wsA.append(exampleA)
for i in range(1, len(colsA) + 1):
    cell = wsA.cell(row=2, column=i)
    cell.font = EX_FONT
    cell.alignment = WRAP
    cell.border = BORDER

# 下拉校验
def dv(ws, options, rng, tip):
    d = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                       allow_blank=True, showDropDown=False)
    d.promptTitle = "请选择"
    d.prompt = tip
    ws.add_data_validation(d)
    d.add(rng)

dv(wsA, ["原创实拍", "翻拍", "混剪", "AI", "达人", "图文"], "D3:D500", "素材类型")
dv(wsA, ["测试中", "已出数据", "已判定"], "L3:L500", "状态(会前必须更新)")
dv(wsA, ["验证成功", "验证失败", "数据不足"], "S3:S500", "结论(不允许空着关行)")

# 条件格式: 状态=测试中 提示未判定; 结论=验证成功 绿、失败 红
wsA.conditional_formatting.add("L3:L500",
    CellIsRule(operator="equal", formula=['"测试中"'],
               fill=PatternFill("solid", fgColor="FFF2CC")))
wsA.conditional_formatting.add("S3:S500",
    CellIsRule(operator="equal", formula=['"验证成功"'],
               fill=PatternFill("solid", fgColor="C6EFCE")))
wsA.conditional_formatting.add("S3:S500",
    CellIsRule(operator="equal", formula=['"验证失败"'],
               fill=PatternFill("solid", fgColor="FFC7CE")))

wsA.freeze_panes = "B2"
wsA.sheet_view.zoomScale = 90

# ============ 表B：结论-验证台账 ============
wsB = wb.create_sheet("台账B-结论验证")
colsB = [
    ("结论编号", 12), ("提出日期", 11), ("来源", 16),
    ("归因结论(一句话,必填)", 32), ("支撑数据", 20), ("置信度", 9),
    ("验证动作", 28), ("负责人", 9), ("验证截止日期", 12),
    ("验证素材ID", 12), ("验证结果", 12), ("后续动作", 10),
    ("状态", 10), ("连续失败次数", 12),
]
wsB.append([c[0] for c in colsB])
for i, (_, w) in enumerate(colsB, 1):
    wsB.column_dimensions[get_column_letter(i)].width = w
    cell = wsB.cell(row=1, column=i)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER

exampleB = [
    "C-2026-001", "2026-10-13", "第1次复盘会",
    "开头出镜人偏老是素材点击率低的原因", "点击率1.2% vs 同梯队2.8%", "中",
    "翻拍同脚本换年轻出镜人,对比点击率", "雪林", "2026-10-20",
    "S-2026-0135", "数据不足", "迭代", "验证中", 0,
]
wsB.append(exampleB)
for i in range(1, len(colsB) + 1):
    cell = wsB.cell(row=2, column=i)
    cell.font = EX_FONT
    cell.alignment = WRAP
    cell.border = BORDER

dv(wsB, ["高", "中", "低"], "F3:F500", "有对照数据=高;单点观察=低")
dv(wsB, ["成立", "不成立", "数据不足"], "K3:K500", "验证结果")
dv(wsB, ["放大", "迭代", "放弃"], "L3:L500", "后续动作")
dv(wsB, ["待验证", "验证中", "已闭环"], "M3:M500", "状态(每次复盘会开头先过此列)")

# 条件格式: 连续失败>=2 红色高亮(南墙检讨); 状态=已闭环 绿; 待验证 黄
wsB.conditional_formatting.add("N3:N500",
    CellIsRule(operator="greaterThanOrEqual", formula=["2"],
               fill=PatternFill("solid", fgColor="FFC7CE")))
wsB.conditional_formatting.add("M3:M500",
    CellIsRule(operator="equal", formula=['"已闭环"'],
               fill=PatternFill("solid", fgColor="C6EFCE")))
wsB.conditional_formatting.add("M3:M500",
    CellIsRule(operator="equal", formula=['"待验证"'],
               fill=PatternFill("solid", fgColor="FFF2CC")))

wsB.freeze_panes = "B2"
wsB.sheet_view.zoomScale = 90

# ============ 表C：填写说明 ============
wsC = wb.create_sheet("填写说明")
wsC.column_dimensions["A"].width = 16
wsC.column_dimensions["B"].width = 90

rowsC = [
    ("【本工作簿用途】", ""),
    ("", ""),
    ("两张台账支撑高乐捷「五段式内容复盘会」(见配套模板文档),使用规则:", ""),
    ("", ""),
    ("台账A-测试标记", "所有上线测试的素材必须先登记后投放;未标记=复盘跟不到单条结果=白测。素材上线当天由投放运营登记,牵头人每周会前核对。"),
    ("台账B-结论验证", "每条复盘归因结论走完「提出→验证→处置」全程,杜绝观点上周提了、这周还没数据反馈。牵头人会上登记,FDE抽查闭环率。"),
    ("", ""),
    ("【哪些单元格要填】", ""),
    ("", ""),
    ("灰色斜体行", "示例行,展示期望填写格式,正式使用时删除(右键行号→删除)。"),
    ("台账A 第2~21列", "字段含义与填写红线详见《高乐捷内容复盘会五段式模板.md》第三节。核心:第5列「测试假设」必填,格式=验证XXX能否带来XXX。"),
    ("台账A 下拉列", "素材类型(D)、状态(L)、结论(S)三列有下拉选项,点击单元格右侧箭头选择。"),
    ("台账B 下拉列", "置信度(F)、验证结果(K)、后续动作(L)、状态(M)四列有下拉。"),
    ("", ""),
    ("【颜色高亮含义】", ""),
    ("", ""),
    ("黄色填充", "台账A状态=测试中(未判定);台账B状态=待验证(超2周升级FDE)。"),
    ("绿色填充", "结论=验证成功(台账A)/已闭环(台账B)。"),
    ("红色填充", "台账A结论=验证失败;台账B连续失败次数>=2,强制触发南墙检讨:同一形式/卖点翻拍3次仍无数据,必须换形式、换人群或暂停。"),
    ("", ""),
    ("【会前检查清单】", ""),
    ("", ""),
    ("1", "台账A所有测试素材状态列已更新(不允许停留在「测试中」超过判定周期3~5天)。"),
    ("2", "台账B「待验证」到期项已逐条过完,未闭环的给新截止时间,连续两次延期升级FDE。"),
    ("3", "在途结论控制在10条以内,超出的先关闭旧的再开新的。"),
    ("", ""),
    ("编制: FDE 马亚鹏 · 2026-10-10 · v1.0", ""),
]
r = 1
for a, b in rowsC:
    wsC.cell(row=r, column=1, value=a)
    wsC.cell(row=r, column=2, value=b)
    ca = wsC.cell(row=r, column=1)
    cb = wsC.cell(row=r, column=2)
    ca.font = Font(name=FONT_NAME, size=11, bold=True) if a.startswith(("【", "台账", "灰色", "1", "2", "3", "编制")) else Font(name=FONT_NAME, size=11)
    cb.font = BODY_FONT
    cb.alignment = WRAP
    if a.startswith("【"):
        ca.fill = SEC_FILL
        wsC.cell(row=r, column=2).fill = SEC_FILL
    r += 1

wb.save(OUT)
print("saved:", OUT)
